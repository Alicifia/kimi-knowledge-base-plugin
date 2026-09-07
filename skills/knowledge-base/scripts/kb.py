#!/usr/bin/env python3
# -*- coding: utf-8 -*-
#
#  kimi-knowledge-base-plugin — 本地知识库引擎
#  Copyright (c) 2026 Alicifia (https://github.com/Alicifia)
#  Licensed under the MIT License. See LICENSE file in the project root
#  for the full license text. 保留版权与许可声明。
#
"""本地知识库引擎（kb.py）

零第三方依赖（仅 Python 标准库），完全离线运行。

功能：
  - init / use   指定并切换知识库本地目录
  - add          导入文档（.md/.markdown/.txt/.epub/.mobi/.azw3，可选 .docx/.pdf）
  - search       向量检索（两级：文档级向量预筛 → 块级向量精排，余弦相似度）
  - related      多文档关联（显式 [[wikilink]]/markdown 链接 + 文档级向量相似）
  - list / remove / status / reindex

存储：知识库根目录下的 .knowledge-base/index.sqlite（文档、块、TF/向量、链接边）。
全局配置（当前激活的知识库）：~/.kimi-work/knowledge-base.json
"""

import argparse
import hashlib
import html as html_mod
import json
import pickle
import re
import sqlite3
import struct
import sys
import unicodedata
import zipfile
from pathlib import Path

CHUNK_TARGET = 500        # 目标块长（字符）
CHUNK_OVERLAP = 80        # 相邻块重叠
DOC_PREFILTER = 8         # 检索时文档级预筛保留的文档数
TEXT_EXT = {".md", ".markdown", ".txt"}
EBOOK_EXT = {".epub", ".mobi", ".azw3", ".azw"}
OPTIONAL_EXT = {".docx", ".pdf"}          # 需 python-docx / pdfplumber
ALL_EXT = TEXT_EXT | EBOOK_EXT | OPTIONAL_EXT
CONFIG_PATH = Path.home() / ".kimi-work" / "knowledge-base.json"
INDEX_DIRNAME = ".knowledge-base"

# ---------------------------------------------------------------- 基础


class ExtractError(Exception):
    """文件可读但解析失败，或缺少所需解析库。"""


def die(msg, code=1):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(code)


def load_config():
    if CONFIG_PATH.exists():
        try:
            return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def save_config(cfg):
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


def require_root(args):
    root = getattr(args, "root", None) or load_config().get("active_root")
    if not root:
        die("尚未指定知识库目录。先运行：kb.py init --root <知识库目录>")
    root = Path(root).expanduser().resolve()
    if not root.exists():
        die(f"知识库目录不存在：{root}")
    return root

# ---------------------------------------------------------------- 数据库

SCHEMA = """
CREATE TABLE IF NOT EXISTS docs(
  id INTEGER PRIMARY KEY,
  path TEXT UNIQUE NOT NULL,      -- 相对知识库根目录的路径
  title TEXT,
  mtime REAL,
  size INTEGER,
  sha TEXT,
  tf BLOB,                        -- pickle({term: count}) 文档级词频
  vec BLOB                        -- pickle({term: weight}) 文档级向量，L2 归一化
);
CREATE TABLE IF NOT EXISTS chunks(
  id INTEGER PRIMARY KEY,
  doc_id INTEGER NOT NULL REFERENCES docs(id) ON DELETE CASCADE,
  ord INTEGER NOT NULL,
  heading TEXT,
  content TEXT NOT NULL,
  tf BLOB,                        -- pickle({term: count})
  vec BLOB                        -- pickle({term: weight})，L2 归一化
);
CREATE TABLE IF NOT EXISTS terms(
  term TEXT PRIMARY KEY,
  df INTEGER NOT NULL
);
CREATE TABLE IF NOT EXISTS links(
  src_doc INTEGER NOT NULL REFERENCES docs(id) ON DELETE CASCADE,
  dst_doc INTEGER REFERENCES docs(id) ON DELETE CASCADE,
  dst_name TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_chunks_doc ON chunks(doc_id);
"""


def open_db(root: Path) -> sqlite3.Connection:
    idx_dir = root / INDEX_DIRNAME
    idx_dir.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(idx_dir / "index.sqlite"))
    conn.execute("PRAGMA foreign_keys=ON")
    conn.executescript(SCHEMA)
    return conn

# ---------------------------------------------------------------- 分词 / 向量化

_WORD_RE = re.compile(r"[a-z0-9_]+")
_CJK_RUN_RE = re.compile(r"[一-鿿㐀-䶿\uf900-\ufaff]+")


def tokenize(text: str):
    """CJK 连续段生成 单字+二元组；拉丁/数字按词。返回 token 列表。"""
    text = unicodedata.normalize("NFKC", text.lower())
    tokens = _WORD_RE.findall(text)
    for run in _CJK_RUN_RE.findall(text):
        if len(run) == 1:
            tokens.append(run)
        else:
            tokens.extend(run)
            tokens.extend(run[i] + run[i + 1] for i in range(len(run) - 1))
    return tokens


def tf_counter(tokens):
    tf = {}
    for t in tokens:
        tf[t] = tf.get(t, 0) + 1
    return tf


def tfidf_vector(tf: dict, idf: dict):
    import math
    vec = {}
    norm = 0.0
    for t, c in tf.items():
        w = (1.0 + math.log(c)) * idf.get(t, 0.0)
        if w > 0:
            vec[t] = w
            norm += w * w
    if norm > 0:
        inv = 1.0 / math.sqrt(norm)
        vec = {t: w * inv for t, w in vec.items()}
    return vec


def cosine(v1: dict, v2: dict) -> float:
    if not v1 or not v2:
        return 0.0
    if len(v1) > len(v2):
        v1, v2 = v2, v1
    return sum(w * v2.get(t, 0.0) for t, w in v1.items())


def load_vec(blob):
    return pickle.loads(blob) if blob else {}

# ---------------------------------------------------------------- 文本抽取

_TAG_BLOCK_RE = re.compile(r"(?is)<(script|style)[^>]*>.*?</\1>")
_TAG_BREAK_RE = re.compile(r"(?i)<\s*(br|/p|/div|/h[1-6]|/li|/tr)[^>]*>")
_TAG_ANY_RE = re.compile(r"(?s)<[^>]+>")


def html_to_text(raw: str) -> str:
    raw = _TAG_BLOCK_RE.sub(" ", raw)
    raw = _TAG_BREAK_RE.sub("\n\n", raw)
    raw = _TAG_ANY_RE.sub(" ", raw)
    raw = html_mod.unescape(raw)
    raw = re.sub(r"[ \t　]+", " ", raw)
    raw = re.sub(r"\n\s*\n+", "\n\n", raw)
    return raw.strip()


def extract_epub(path: Path) -> str:
    parts = []
    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist()
                 if n.lower().endswith((".xhtml", ".html", ".htm"))
                 and not n.startswith("__MACOSX")]
        if not names:
            raise ExtractError("epub 内没有 xhtml/html 内容文件")
        for name in sorted(names):
            try:
                raw = zf.read(name).decode("utf-8", errors="replace")
            except Exception:
                continue
            text = html_to_text(raw)
            if text:
                parts.append(text)
    if not parts:
        raise ExtractError("epub 内容抽取后为空")
    return "\n\n".join(parts)


def _palmdoc_decompress(data: bytes) -> bytes:
    """PalmDOC LZ77 解压。"""
    out = bytearray()
    i, n = 0, len(data)
    while i < n:
        c = data[i]
        i += 1
        if c == 0:
            out.append(0)
        elif c <= 8:                       # 1..8: 原文复制 c 字节
            out += data[i:i + c]
            i += c
        elif c < 0x80:
            out.append(c)
        elif c >= 0xC0:                    # 空格 + 字符
            out += b" " + bytes([c ^ 0x80])
        else:                              # 0x80..0xBF: LZ77 回引
            if i >= n:
                break
            c = (c << 8) | data[i]
            i += 1
            dist = (c & 0x3FFF) >> 3
            length = (c & 0x07) + 3
            if dist == 0 or dist > len(out):
                continue
            for _ in range(length):
                out.append(out[-dist])
    return bytes(out)


def extract_mobi(path: Path) -> str:
    data = path.read_bytes()
    if len(data) < 78:
        raise ExtractError("文件太小，不是有效 mobi/azw3/azw")
    num_records = struct.unpack(">H", data[76:78])[0]
    if num_records < 2:
        raise ExtractError("mobi 记录数异常")
    offsets = []
    for i in range(num_records):
        off = struct.unpack(">I", data[78 + i * 8:82 + i * 8])[0]
        offsets.append(off)
    offsets.append(len(data))

    def section_text(start):
        """从第 start 条记录处的 PalmDOC/MOBI 头抽取一个区段的正文。"""
        rec0 = data[offsets[start]:offsets[start + 1]]
        if len(rec0) < 32 or rec0[16:20] != b"MOBI":
            return None, None
        compression = struct.unpack(">H", rec0[0:2])[0]
        if compression not in (1, 2):
            return None, compression
        text_length = struct.unpack(">I", rec0[4:8])[0]
        record_count = struct.unpack(">H", rec0[8:10])[0]
        encoding = struct.unpack(">I", rec0[28:32])[0]
        codec = "utf-8" if encoding == 65001 else "cp1252"
        parts = []
        for i in range(start + 1, min(start + record_count, num_records - 1) + 1):
            rec = data[offsets[i]:offsets[i + 1]]
            if compression == 2:
                rec = _palmdoc_decompress(rec)
            parts.append(rec)
        raw = b"".join(parts)[:text_length]
        return html_to_text(raw.decode(codec, errors="replace")), compression

    # 找所有区段头（双区段文件里 KF8 区段通常在文件后部），取正文最长者
    best = ""
    huff_seen = None
    found_header = False
    for i in range(num_records):
        if data[offsets[i] + 16:offsets[i] + 20] != b"MOBI":
            continue
        found_header = True
        text, comp = section_text(i)
        if text is None:
            huff_seen = comp
            continue
        if len(text) > len(best):
            best = text
    if not found_header:
        raise ExtractError("找不到 MOBI 头，不是有效 mobi/azw3/azw")
    if len(best) < 100:
        if huff_seen is not None:
            raise ExtractError(f"HuffCDIC 压缩（类型 {huff_seen}），暂不支持")
        raise ExtractError("mobi 正文抽取后过短，疑似解析失败")
    return best


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".epub":
            return extract_epub(path)
        if ext in (".mobi", ".azw3", ".azw"):
            return extract_mobi(path)
        if ext == ".docx":
            try:
                from docx import Document
            except ImportError:
                raise ExtractError("缺少 python-docx，无法解析 .docx")
            d = Document(str(path))
            return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
        if ext == ".pdf":
            try:
                import pdfplumber
            except ImportError:
                raise ExtractError("缺少 pdfplumber，无法解析 .pdf")
            with pdfplumber.open(str(path)) as pdf:
                return "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except ExtractError:
        raise
    except Exception as e:
        raise ExtractError(f"{type(e).__name__}: {e}") from e
    raise ExtractError(f"不支持的格式：{ext}")

# ---------------------------------------------------------------- 切块

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def chunk_markdown(text: str):
    """按标题层级切块，超长块按段落再拆。返回 [(heading_path, content)]"""
    lines = text.splitlines()
    sections = []
    heading_stack = []
    buf = []

    def flush():
        nonlocal buf
        body = "\n".join(buf).strip()
        buf = []
        if body:
            sections.append((" / ".join(t for _, t in heading_stack), body))

    for line in lines:
        m = _HEADING_RE.match(line)
        if m:
            flush()
            level = len(m.group(1))
            heading_stack = [(l, t) for l, t in heading_stack if l < level]
            heading_stack.append((level, m.group(2).strip()))
        else:
            buf.append(line)
    flush()

    chunks = []
    for heading, body in sections:
        paras = re.split(r"\n\s*\n", body)
        cur = ""
        for p in paras:
            p = p.strip()
            if not p:
                continue
            if cur and len(cur) + len(p) > CHUNK_TARGET:
                chunks.append((heading, cur))
                cur = cur[-CHUNK_OVERLAP:] + "\n" + p if CHUNK_OVERLAP < len(cur) else p
            else:
                cur = cur + "\n\n" + p if cur else p
            while len(cur) > CHUNK_TARGET * 2:
                chunks.append((heading, cur[:CHUNK_TARGET]))
                cur = cur[CHUNK_TARGET - CHUNK_OVERLAP:]
        if cur.strip():
            chunks.append((heading, cur))
    return [(h, c) for h, c in chunks if c.strip()]

# ---------------------------------------------------------------- 链接解析

_WIKILINK_RE = re.compile(r"\[\[([^\[\]|]+)(?:\|[^\[\]]*)?\]\]")
_MDLINK_RE = re.compile(r"\[[^\]]*\]\(([^)]+\.(?:md|markdown|txt|docx|pdf|epub|mobi|azw3))[^)]*\)", re.I)


def extract_links(text: str):
    targets = set()
    for m in _WIKILINK_RE.finditer(text):
        targets.add(m.group(1).strip())
    for m in _MDLINK_RE.finditer(text):
        targets.add(Path(m.group(1).strip()).stem)
    return targets

# ---------------------------------------------------------------- 索引


def iter_files(paths):
    for raw in paths:
        p = Path(raw).expanduser()
        if not p.is_absolute():
            p = Path.cwd() / p
        if p.is_dir():
            for f in sorted(p.rglob("*")):
                if f.is_file() and f.suffix.lower() in ALL_EXT:
                    if INDEX_DIRNAME not in f.parts:
                        yield f
        elif p.is_file():
            yield p
        else:
            print(f"WARN: 路径不存在，跳过：{raw}", file=sys.stderr)


def rebuild_idf(conn):
    """由库存 TF 重算全库 idf，只回写文档级向量（块级向量检索时按需现算，不重新分词）。"""
    import math
    n = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    df = {t: d for t, d in conn.execute("SELECT term, df FROM terms")}
    idf = {t: math.log((1 + n) / (1 + d)) + 1.0 for t, d in df.items()}
    for doc_id, tf_blob in conn.execute("SELECT id, tf FROM docs WHERE tf IS NOT NULL"):
        vec = tfidf_vector(load_vec(tf_blob), idf)
        conn.execute("UPDATE docs SET vec=? WHERE id=?", (pickle.dumps(vec), doc_id))
    conn.commit()


def apply_terms(conn, tf_counters, sign):
    """批量更新词表 df。先聚合成 {term: delta}，再 executemany 一次落库。"""
    agg = {}
    for tf in tf_counters:
        for t in tf:
            agg[t] = agg.get(t, 0) + sign
    ins_val = max(sign, 1)
    conn.executemany(
        "INSERT INTO terms(term, df) VALUES(?, ?) "
        "ON CONFLICT(term) DO UPDATE SET df = MAX(df + ?, 0)",
        [(t, ins_val, sign) for t in agg],
    )


def index_file(conn, root: Path, f: Path, force=False):
    """返回 (status, detail)。status ∈ indexed/unchanged/empty/failed"""
    try:
        rel = str(f.resolve().relative_to(root.resolve()))
    except ValueError:
        die(f"文件不在知识库目录内：{f}\n知识库根目录是 {root}，请把文件移进去或用对应库的根目录。")
    st = f.stat()
    sha = hashlib.sha1(f.read_bytes()).hexdigest()
    row = conn.execute("SELECT id, sha FROM docs WHERE path=?", (rel,)).fetchone()
    if row and row[1] == sha and not force:
        return "unchanged", ""

    try:
        text = extract_text(f)
    except ExtractError as e:
        return "failed", str(e)
    if not text.strip():
        return "empty", ""

    if row:
        old_tfs = [load_vec(r[0]) for r in
                   conn.execute("SELECT tf FROM chunks WHERE doc_id=?", (row[0],))]
        apply_terms(conn, old_tfs, sign=-1)
        conn.execute("DELETE FROM docs WHERE id=?", (row[0],))

    title = f.stem
    m = re.search(r"^#\s+(.+)$", text, re.M)
    if m:
        title = m.group(1).strip()
    cur = conn.execute(
        "INSERT INTO docs(path, title, mtime, size, sha) VALUES(?,?,?,?,?)",
        (rel, title, st.st_mtime, st.st_size, sha),
    )
    doc_id = cur.lastrowid

    chunks = chunk_markdown(text)
    tfs = []
    doc_tf = {}
    rows = []
    for i, (heading, content) in enumerate(chunks):
        tf = tf_counter(tokenize(content))
        tfs.append(tf)
        for t, c in tf.items():
            doc_tf[t] = doc_tf.get(t, 0) + c
        rows.append((doc_id, i, heading, content, pickle.dumps(tf)))
    conn.executemany(
        "INSERT INTO chunks(doc_id, ord, heading, content, tf) VALUES(?,?,?,?,?)",
        rows,
    )
    apply_terms(conn, tfs, sign=+1)
    conn.execute("UPDATE docs SET tf=? WHERE id=?", (pickle.dumps(doc_tf), doc_id))
    for target in extract_links(text):
        dst = conn.execute(
            "SELECT id FROM docs WHERE title=? OR path LIKE ?", (target, f"%{target}.%")
        ).fetchone()
        conn.execute(
            "INSERT INTO links(src_doc, dst_doc, dst_name) VALUES(?,?,?)",
            (doc_id, dst[0] if dst else None, target),
        )
    return "indexed", f"{len(chunks)} 块"

# ---------------------------------------------------------------- 命令


def cmd_init(args):
    root = Path(args.root).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    conn = open_db(root)
    conn.close()
    cfg = load_config()
    cfg["active_root"] = str(root)
    save_config(cfg)
    print(f"OK: 知识库已就绪 -> {root}")
    print(f"索引位置 -> {root / INDEX_DIRNAME / 'index.sqlite'}")
    print("下一步：kb.py add <文件或目录> 导入文档")


def cmd_use(args):
    root = Path(args.root).expanduser().resolve()
    if not (root / INDEX_DIRNAME / "index.sqlite").exists():
        die(f"{root} 还不是知识库（缺少索引）。先运行：kb.py init --root \"{root}\"")
    cfg = load_config()
    cfg["active_root"] = str(root)
    save_config(cfg)
    print(f"OK: 已切换到知识库 -> {root}")


def cmd_add(args):
    root = require_root(args)
    conn = open_db(root)
    stats = {"indexed": 0, "unchanged": 0, "empty": 0, "failed": 0}
    for f in iter_files(args.paths):
        status, detail = index_file(conn, root, f, force=args.force)
        stats[status] += 1
        if status == "indexed":
            print(f"  + {f.name}（{detail}）")
            conn.commit()
        elif status == "failed":
            print(f"  ✗ 读取失败 {f.name}: {detail}", file=sys.stderr)
            if args.strict:
                conn.close()
                die(f"按 --strict 要求退出：文件读取失败 {f}（{detail}）", code=2)
    conn.commit()
    if stats["indexed"]:
        print("正在重建向量索引……", file=sys.stderr)
        rebuild_idf(conn)
    conn.close()
    print(f"OK: 新入库 {stats['indexed']}，未变化 {stats['unchanged']}，"
          f"空文件 {stats['empty']}，失败 {stats['failed']}")


def cmd_remove(args):
    root = require_root(args)
    conn = open_db(root)
    removed = 0
    for raw in args.paths:
        rel = None
        p = Path(raw).expanduser()
        if p.is_absolute():
            try:
                rel = str(p.resolve().relative_to(root.resolve()))
            except ValueError:
                pass
        row = conn.execute(
            "SELECT id FROM docs WHERE path=? OR path=? OR title=?",
            (rel or raw, raw, Path(raw).stem),
        ).fetchone()
        if row:
            tfs = [load_vec(r[0]) for r in
                   conn.execute("SELECT tf FROM chunks WHERE doc_id=?", (row[0],))]
            apply_terms(conn, tfs, sign=-1)
            conn.execute("DELETE FROM docs WHERE id=?", (row[0],))
            removed += 1
            print(f"  - {raw}")
    conn.commit()
    if removed:
        rebuild_idf(conn)
    conn.close()
    print(f"OK: 已移除 {removed} 篇文档")


def _idf_map(conn):
    import math
    n = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    df = {t: d for t, d in conn.execute("SELECT term, df FROM terms")}
    return {t: math.log((1 + n) / (1 + d)) + 1.0 for t, d in df.items()}, n


def cmd_search(args):
    root = require_root(args)
    conn = open_db(root)
    idf, n = _idf_map(conn)
    if n == 0:
        die("知识库是空的。先运行：kb.py add <文件或目录>")
    qvec = tfidf_vector(tf_counter(tokenize(args.query)), idf)
    if not qvec:
        die("查询分词后为空，换个说法试试。")

    # 第一级：文档级向量预筛
    doc_scores = []
    for doc_id, vec_blob in conn.execute("SELECT id, vec FROM docs WHERE vec IS NOT NULL"):
        s = cosine(qvec, load_vec(vec_blob))
        if s > 0:
            doc_scores.append((s, doc_id))
    doc_scores.sort(key=lambda x: -x[0])
    top_docs = doc_scores[:DOC_PREFILTER]
    if not top_docs:
        print(json.dumps({"kb_root": str(root), "query": args.query, "hits": []},
                         ensure_ascii=False, indent=2))
        conn.close()
        return
    doc_score_map = {doc_id: s for s, doc_id in top_docs}
    placeholders = ",".join("?" * len(top_docs))

    # 第二级：候选文档内块级精排（块向量由 TF 按需现算）
    scored = []
    for cid, doc_id, heading, content, tf_blob in conn.execute(
            f"SELECT id, doc_id, heading, content, tf FROM chunks WHERE doc_id IN ({placeholders})",
            tuple(doc_score_map)):
        s = cosine(qvec, tfidf_vector(load_vec(tf_blob), idf))
        if s > 0:
            scored.append((s, cid, doc_id, heading, content))
    scored.sort(key=lambda x: -x[0])

    seen_docs = {}
    out = []
    for s, cid, doc_id, heading, content in scored:
        if args.per_doc and doc_id in seen_docs:
            continue
        seen_docs[doc_id] = True
        path, title = conn.execute("SELECT path, title FROM docs WHERE id=?", (doc_id,)).fetchone()
        out.append({
            "score": round(s, 4),
            "doc_score": round(doc_score_map[doc_id], 4),
            "path": str(root / path),
            "rel_path": path,
            "title": title,
            "heading": heading or "",
            "snippet": " ".join(content.split())[:220],
        })
        if len(out) >= args.top:
            break
    print(json.dumps({"kb_root": str(root), "query": args.query, "hits": out},
                     ensure_ascii=False, indent=2))
    conn.close()


def cmd_related(args):
    root = require_root(args)
    conn = open_db(root)
    target = args.path
    row = conn.execute(
        "SELECT id, path, title, vec FROM docs WHERE path=? OR title=? OR path LIKE ?",
        (target, Path(target).stem, f"%{target}%"),
    ).fetchone()
    if not row:
        die(f"库中找不到文档：{target}")
    doc_id, path, title, my_vec_blob = row
    my_vec = load_vec(my_vec_blob)

    # 1) 显式链接（双向）
    linked = {}
    for dst, name in conn.execute("SELECT dst_doc, dst_name FROM links WHERE src_doc=?", (doc_id,)):
        if dst:
            p, t = conn.execute("SELECT path, title FROM docs WHERE id=?", (dst,)).fetchone()
            linked[dst] = (p, t)
        else:
            print(f"  (外链未入库: {name})", file=sys.stderr)
    for src, in conn.execute("SELECT src_doc FROM links WHERE dst_doc=?", (doc_id,)):
        if src != doc_id:
            p, t = conn.execute("SELECT path, title FROM docs WHERE id=?", (src,)).fetchone()
            linked.setdefault(src, (p, t))

    # 2) 文档级向量相似
    sims = []
    for oid, vec_blob in conn.execute(
            "SELECT id, vec FROM docs WHERE id != ? AND vec IS NOT NULL", (doc_id,)):
        s = cosine(my_vec, load_vec(vec_blob))
        if s > 0:
            sims.append((s, oid))
    sims.sort(key=lambda x: -x[0])
    similar = []
    for s, oid in sims[:args.top]:
        p, t = conn.execute("SELECT path, title FROM docs WHERE id=?", (oid,)).fetchone()
        similar.append({"score": round(s, 4), "path": str(root / p), "title": t,
                        "explicit_link": oid in linked})

    print(json.dumps({
        "doc": {"path": str(root / path), "title": title},
        "explicit_links": [{"path": str(root / p), "title": t} for p, t in linked.values()],
        "similar_docs": similar,
    }, ensure_ascii=False, indent=2))
    conn.close()


def cmd_list(args):
    root = require_root(args)
    conn = open_db(root)
    rows = conn.execute(
        "SELECT d.path, d.title, COUNT(c.id) FROM docs d "
        "LEFT JOIN chunks c ON c.doc_id=d.id GROUP BY d.id ORDER BY d.path"
    ).fetchall()
    for path, title, nchunks in rows:
        print(f"{path}\t{title}\t{nchunks} 块")
    print(f"共 {len(rows)} 篇文档")
    conn.close()


def cmd_status(args):
    cfg = load_config()
    root = cfg.get("active_root")
    if not root:
        print("尚未指定知识库目录。运行：kb.py init --root <目录>")
        return
    root = Path(root)
    conn = open_db(root)
    docs = conn.execute("SELECT COUNT(*) FROM docs").fetchone()[0]
    chunks = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    terms = conn.execute("SELECT COUNT(*) FROM terms").fetchone()[0]
    links = conn.execute("SELECT COUNT(*) FROM links").fetchone()[0]
    print(json.dumps({
        "active_root": str(root),
        "docs": docs, "chunks": chunks, "vocab_terms": terms, "links": links,
    }, ensure_ascii=False, indent=2))
    conn.close()


def cmd_reindex(args):
    root = require_root(args)
    conn = open_db(root)
    conn.execute("DELETE FROM links")
    conn.execute("DELETE FROM chunks")
    conn.execute("DELETE FROM docs")
    conn.execute("DELETE FROM terms")
    conn.commit()
    stats = {"indexed": 0, "failed": 0}
    for f in sorted(root.rglob("*")):
        if f.is_file() and f.suffix.lower() in ALL_EXT and INDEX_DIRNAME not in f.parts:
            status, detail = index_file(conn, root, f, force=True)
            if status == "indexed":
                stats["indexed"] += 1
                conn.commit()
            elif status == "failed":
                stats["failed"] += 1
                print(f"  ✗ 读取失败 {f.name}: {detail}", file=sys.stderr)
                if getattr(args, "strict", False):
                    conn.close()
                    die(f"按 --strict 要求退出：文件读取失败 {f}（{detail}）", code=2)
    conn.commit()
    if stats["indexed"]:
        print("正在重建向量索引……", file=sys.stderr)
        rebuild_idf(conn)
    conn.close()
    print(f"OK: 全量重建完成，共 {stats['indexed']} 篇文档，失败 {stats['failed']}")

# ---------------------------------------------------------------- 入口


def main():
    ap = argparse.ArgumentParser(prog="kb.py", description="本地知识库引擎")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("init", help="指定知识库目录并初始化")
    p.add_argument("--root", required=True)
    p.set_defaults(fn=cmd_init)

    p = sub.add_parser("use", help="切换当前激活的知识库")
    p.add_argument("--root", required=True)
    p.set_defaults(fn=cmd_use)

    p = sub.add_parser("add", help="导入文件或目录")
    p.add_argument("paths", nargs="+")
    p.add_argument("--root", default=None)
    p.add_argument("--force", action="store_true")
    p.add_argument("--strict", action="store_true", help="任何文件读取失败立即退出（exit 2）")
    p.set_defaults(fn=cmd_add)

    p = sub.add_parser("remove", help="从库中移除文档")
    p.add_argument("paths", nargs="+")
    p.add_argument("--root", default=None)
    p.set_defaults(fn=cmd_remove)

    p = sub.add_parser("search", help="向量检索")
    p.add_argument("query")
    p.add_argument("--top", type=int, default=8)
    p.add_argument("--per-doc", action="store_true", help="每篇文档只保留最佳块")
    p.add_argument("--root", default=None)
    p.set_defaults(fn=cmd_search)

    p = sub.add_parser("related", help="查某篇文档的关联文档")
    p.add_argument("path")
    p.add_argument("--top", type=int, default=5)
    p.add_argument("--root", default=None)
    p.set_defaults(fn=cmd_related)

    p = sub.add_parser("list", help="列出库内文档")
    p.add_argument("--root", default=None)
    p.set_defaults(fn=cmd_list)

    p = sub.add_parser("status", help="查看当前知识库状态")
    p.set_defaults(fn=cmd_status)

    p = sub.add_parser("reindex", help="全量重建索引")
    p.add_argument("--root", default=None)
    p.add_argument("--strict", action="store_true")
    p.set_defaults(fn=cmd_reindex)

    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
